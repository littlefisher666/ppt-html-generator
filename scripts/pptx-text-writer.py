#!/usr/bin/env python3
"""Shared python-pptx text writer for the P3 hybrid deck.

Text specs come from P3/text-plan.json (built from P3/text-nodes.json).
Only the writer primitives live here; deck assembly is in
apply-hybrid-text-plan.py and probe calibration in calibrate-hybrid-text.py.
"""

from __future__ import annotations

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU_PER_PX = 7620


def px_to_emu(value: float) -> int:
    return int(round(float(value) * EMU_PER_PX))


def emu(value: float) -> Emu:
    return Emu(px_to_emu(value))


def hex_rgb(value: str) -> RGBColor:
    text = str(value or "").strip().lstrip("#")
    if len(text) != 6:
        raise ValueError(f"invalid color: {value}")
    return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


def zero_text_frame_margins(tf) -> None:
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)


def disable_autofit(tf) -> None:
    try:
        body_pr = tf._txBody.bodyPr
        for child in list(body_pr):
            if child.tag.endswith("spAutoFit") or child.tag.endswith("normAutofit"):
                body_pr.remove(child)
    except Exception:
        pass


def set_run_font(run, name, size_pt, bold, color, char_spacing_px=None) -> None:
    run.font.name = name
    run.font.size = Pt(float(size_pt))
    run.font.bold = bool(bold)
    run.font.color.rgb = hex_rgb(color)
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", name)
    latin = r_pr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(r_pr, qn("a:latin"))
    latin.set("typeface", name)
    r_pr.set("b", "1" if bold else "0")
    if char_spacing_px is not None:
        try:
            spc_val = int(round(float(char_spacing_px) * 0.6 * 100))
        except (TypeError, ValueError):
            spc_val = 0
        if spc_val == 0:
            if "spc" in r_pr.attrib:
                del r_pr.attrib["spc"]
        else:
            r_pr.set("spc", str(spc_val))


def set_line_spacing(paragraph, style: dict) -> None:
    """Write line spacing from style.

    Prefer exact spacing (a:spcPts) when lineSpacingPx is present: LibreOffice and
    PowerPoint treat spcPct as a multiple of the font's natural line height, which
    inflates CSS line-heights by ~1.2x. Exact pt spacing reproduces CSS line boxes.
    """
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag.endswith("lnSpc"):
            pPr.remove(child)
    exact_px = style.get("lineSpacingPx")
    if exact_px is not None and float(exact_px) > 0:
        ln_spc = etree.SubElement(pPr, qn("a:lnSpc"))
        spc_pts = etree.SubElement(ln_spc, qn("a:spcPts"))
        spc_pts.set("val", str(int(round(float(exact_px) * 0.6 * 100))))
        return
    ratio = style.get("lineHeightRatio")
    if ratio is None:
        return
    ln_spc = etree.SubElement(pPr, qn("a:lnSpc"))
    spc_pct = etree.SubElement(ln_spc, qn("a:spcPct"))
    spc_pct.set("val", str(int(round(float(ratio) * 100000))))


def split_text_lines(spec: dict) -> list[str]:
    write = spec.get("write") or {}
    planned = write.get("linesForWrite")
    if isinstance(planned, list) and planned:
        return [str(item) for item in planned]
    text = str(spec.get("text") or "")
    if write.get("useHardLineBreaks") and "\n" in text:
        return text.split("\n")
    return [text]


def _clear_paragraph_content(paragraph) -> None:
    """Remove runs/breaks while keeping paragraph properties."""
    p = paragraph._p
    for child in list(p):
        tag = child.tag
        if tag.endswith("}r") or tag.endswith("}br") or tag.endswith("}fld") or tag.endswith("}endParaRPr"):
            p.remove(child)


def _apply_end_para_rpr(paragraph, style: dict) -> None:
    try:
        end = paragraph._p.get_or_add_endParaRPr()
        end.set("sz", str(int(round(float(style["fontSizePt"]) * 100))))
        end.set("b", "1" if style.get("bold", False) else "0")
        ea = end.find(qn("a:ea"))
        if ea is None:
            ea = etree.SubElement(end, qn("a:ea"))
        ea.set("typeface", style["fontName"])
        latin = end.find(qn("a:latin"))
        if latin is None:
            latin = etree.SubElement(end, qn("a:latin"))
        latin.set("typeface", style["fontName"])
    except Exception:
        pass


def _add_line_break(paragraph, style: dict) -> None:
    br = etree.SubElement(paragraph._p, qn("a:br"))
    r_pr = etree.SubElement(br, qn("a:rPr"))
    r_pr.set("sz", str(int(round(float(style["fontSizePt"]) * 100))))
    r_pr.set("b", "1" if style.get("bold", False) else "0")
    ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", style["fontName"])
    latin = etree.SubElement(r_pr, qn("a:latin"))
    latin.set("typeface", style["fontName"])


def fill_paragraph(paragraph, text: str, style: dict) -> None:
    paragraph.alignment = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }.get(style.get("textAlign") or "left", PP_ALIGN.LEFT)
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    set_line_spacing(paragraph, style)
    _clear_paragraph_content(paragraph)
    run = paragraph.add_run()
    run.text = text
    set_run_font(
        run,
        style["fontName"],
        style["fontSizePt"],
        style.get("bold", False),
        style["color"],
        char_spacing_px=style.get("letterSpacingPx"),
    )
    # Sync end-para defaults so some renderers do not fall back to default size.
    _apply_end_para_rpr(paragraph, style)


def fill_paragraph_with_hard_breaks(paragraph, lines: list[str], style: dict) -> None:
    """Write multi-line text as one paragraph with a:br hard breaks.

    Multiple paragraphs introduce extra paragraph spacing risk in PPT.
    """
    paragraph.alignment = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }.get(style.get("textAlign") or "left", PP_ALIGN.LEFT)
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    set_line_spacing(paragraph, style)
    _clear_paragraph_content(paragraph)

    safe_lines = lines if lines else [""]
    for index, line in enumerate(safe_lines):
        if index > 0:
            _add_line_break(paragraph, style)
        run = paragraph.add_run()
        run.text = str(line)
        set_run_font(
            run,
            style["fontName"],
            style["fontSizePt"],
            style.get("bold", False),
            style["color"],
            char_spacing_px=style.get("letterSpacingPx"),
        )
    _apply_end_para_rpr(paragraph, style)


def add_textbox(slide, spec: dict):
    style = spec["style"]
    box = spec["box"]
    shape = slide.shapes.add_textbox(emu(box["x"]), emu(box["y"]), emu(box["width"]), emu(box["height"]))
    shape.name = spec["shapeName"]
    tf = shape.text_frame
    write = spec.get("write") or {}
    tf.word_wrap = bool(write.get("wordWrap", False))
    zero_text_frame_margins(tf)
    disable_autofit(tf)
    # 明确禁止自动缩字，溢出应通过盒尺寸/硬换行解决
    try:
        body_pr = tf._txBody.bodyPr
        for child in list(body_pr):
            if child.tag.endswith("spAutoFit") or child.tag.endswith("normAutofit"):
                body_pr.remove(child)
        etree.SubElement(body_pr, qn("a:noAutofit"))
    except Exception:
        pass
    anchor_map = {"bottom": "b", "middle": "ctr", "top": "t"}
    try:
        tf._txBody.bodyPr.set("anchor", anchor_map.get(style.get("verticalAlign") or "top", "t"))
    except Exception:
        pass

    lines = split_text_lines(spec)
    if not lines:
        lines = [""]
    first = tf.paragraphs[0]
    # Prefer single paragraph + hard breaks to avoid multi-paragraph spacing inflation.
    if len(lines) > 1 or bool(write.get("useHardLineBreaks")):
        fill_paragraph_with_hard_breaks(first, lines, style)
    else:
        fill_paragraph(first, lines[0], style)
    return shape


