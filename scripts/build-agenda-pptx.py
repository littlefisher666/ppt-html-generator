#!/usr/bin/env python3
"""从 agenda-template.json 生成目录页 PPTX。

固定层使用经 manifest 哈希验证的无 Logo 高清整图（背景+CONTENTS+目录+中部图形），
彩色 Logo 从已验证的 deck 品牌资源独立放置；右侧目录项长条背景与编号方块使用独立内置高清资产；
仅编号与标题为可编辑文本。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt


EMU_PER_PX = 7620


def px(value):
    return int(round(float(value) * EMU_PER_PX))


def hex_rgb(value):
    text = value.strip().lstrip("#")
    if len(text) != 6:
        raise ValueError(f"不支持的颜色：{value}")
    return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


def zero_text_frame_margins(tf):
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")
    except Exception:
        pass


def disable_autofit(tf):
    try:
        body_pr = tf._txBody.bodyPr
        for child in list(body_pr):
            if child.tag.endswith("spAutoFit") or child.tag.endswith("normAutofit"):
                body_pr.remove(child)
        body_pr.set("anchor", "ctr")
    except Exception:
        pass


def set_run_font(run, name, size_pt, bold, color):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = hex_rgb(color)
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", name)
    latin = r_pr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(r_pr, qn("a:latin"))
    latin.set("typeface", name)
    if not bold:
        r_pr.set("b", "0")


def add_textbox(slide, x, y, w, h, text, font_name, font_pt, color, align_center=False):
    box = slide.shapes.add_textbox(Emu(px(x)), Emu(px(y)), Emu(px(w)), Emu(px(h)))
    tf = box.text_frame
    tf.word_wrap = not align_center
    zero_text_frame_margins(tf)
    disable_autofit(tf)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if align_center else PP_ALIGN.LEFT
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    if p.runs:
        p.runs[0].text = ""
    run = p.add_run()
    run.text = str(text)
    set_run_font(run, font_name, float(font_pt), False, color)
    if not align_center:
        try:
            pPr = p._p.get_or_add_pPr()
            lnSpc = pPr.find(qn("a:lnSpc"))
            if lnSpc is not None:
                pPr.remove(lnSpc)
            lnSpc = etree.SubElement(pPr, qn("a:lnSpc"))
            spcPct = etree.SubElement(lnSpc, qn("a:spcPct"))
            spcPct.set("val", "130000")
        except Exception:
            pass
    return box


def resolve_asset_path(model, model_path: Path, asset: dict, label: str) -> Path:
    candidates = []
    if asset.get("path"):
        candidates.append(Path(asset["path"]))
    if asset.get("relativePath"):
        candidates.append(model_path.parent / asset["relativePath"])
        candidates.append(Path(model.get("deckRoot", "")) / asset["relativePath"])
        candidates.append(
            Path(model.get("deckRoot", "")) / "exports" / "pptx" / asset["relativePath"]
        )
    for path in candidates:
        if path and path.is_file():
            return path.resolve()
    raise FileNotFoundError(f"未找到{label}：{asset}")


def resolve_item_chrome(model, model_path: Path, item_height) -> dict:
    chromes = model.get("itemChromes") or {}
    key = str(item_height)
    # 兼容 float key 的字符串化差异
    chrome = chromes.get(key)
    if chrome is None:
        for cand_key, cand in chromes.items():
            try:
                if abs(float(cand_key) - float(item_height)) < 0.01:
                    chrome = cand
                    break
            except Exception:
                continue
    if chrome is None:
        raise FileNotFoundError(
            f"未找到 itemHeight={item_height} 的目录条目组件；请先运行 build-agenda-template.mjs"
        )
    bar = chrome.get("bar")
    number_block = chrome.get("numberBlock")
    if not bar or not number_block:
        raise ValueError(
            f"itemHeight={item_height} 的目录条目组件不是分离的长条/编号方块资产；请重新运行 build-agenda-template.mjs"
        )
    return {
        **chrome,
        "bar": {**bar, "resolvedPath": resolve_asset_path(model, model_path, bar, f"条目长条背景({item_height})")},
        "numberBlock": {**number_block, "resolvedPath": resolve_asset_path(model, model_path, number_block, f"编号方块({item_height})")},
    }


def add_item(slide, item, metrics, colors, chrome):
    """独立高清长条/编号方块 + 可编辑编号/标题。"""
    bar = chrome["bar"]
    bar_offset = bar.get("placeOffset") or {"x": 0, "y": 0}
    bar_logical = bar.get("logical") or {"width": item["width"], "height": item["height"]}
    bar_x = float(item["x"]) + float(bar_offset.get("x", 0))
    bar_y = float(item["y"]) + float(bar_offset.get("y", 0))

    slide.shapes.add_picture(
        str(bar["resolvedPath"]),
        Emu(px(bar_x)),
        Emu(px(bar_y)),
        width=Emu(px(float(bar_logical.get("width", item["width"])))),
        height=Emu(px(float(bar_logical.get("height", item["height"])))),
    )

    number_block = chrome["numberBlock"]
    number_offset = number_block.get("placeOffset") or {"x": 0, "y": 0}
    number_logical = number_block.get("logical") or {
        "width": item["numberBox"]["width"],
        "height": item["numberBox"]["height"],
    }
    slide.shapes.add_picture(
        str(number_block["resolvedPath"]),
        Emu(px(float(item["x"]) + float(number_offset.get("x", 0)))),
        Emu(px(float(item["y"]) + float(number_offset.get("y", 0)))),
        width=Emu(px(float(number_logical.get("width", item["numberBox"]["width"])))),
        height=Emu(px(float(number_logical.get("height", item["numberBox"]["height"])))),
    )

    nb = item["numberBox"]
    add_textbox(
        slide,
        nb["x"],
        nb["y"],
        nb["width"],
        nb["height"],
        item["number"],
        colors["fontFamily"],
        metrics["numberFontPt"],
        colors["numberColor"],
        align_center=True,
    )

    tb = item["titleBox"]
    add_textbox(
        slide,
        tb["x"],
        tb["y"],
        tb["width"],
        tb["height"],
        item["title"],
        colors["fontFamily"],
        metrics["titleFontPt"],
        colors["titleColor"],
        align_center=False,
    )


def build_pptx(model, model_path: Path, output_path: Path) -> None:
    pages = model["pages"]
    canvas = model["template"]["canvas"]
    fixed_image_path = resolve_asset_path(
        model, model_path, model.get("fixedImage") or {}, "目录固定高清底图 fixedImage"
    )
    logo = model.get("brandLogo") or {}
    logo_path = resolve_asset_path(model, model_path, logo, "目录页彩色 Logo")
    logical = (model.get("fixedImage") or {}).get("logical") or {
        "x": 0,
        "y": 0,
        "width": canvas["width"],
        "height": canvas["height"],
    }

    prs = Presentation()
    prs.slide_width = Emu(px(canvas["width"]))
    prs.slide_height = Emu(px(canvas["height"]))
    blank = prs.slide_layouts[6]

    colors = {
        "titleColor": model["template"]["itemBase"]["titleColor"],
        "numberColor": model["template"]["itemBase"]["numberColor"],
        "fontFamily": model["template"]["itemBase"]["fontFamily"],
    }

    for page in pages:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(fixed_image_path),
            Emu(px(logical["x"])),
            Emu(px(logical["y"])),
            width=Emu(px(logical["width"])),
            height=Emu(px(logical["height"])),
        )
        slide.shapes.add_picture(
            str(logo_path),
            Emu(px(logo["x"])),
            Emu(px(logo["y"])),
            width=Emu(px(logo["width"])),
            height=Emu(px(logo["height"])),
        )

        layout = page["layout"]
        metrics = layout["metrics"]
        chrome = resolve_item_chrome(model, model_path, metrics["itemHeight"])
        for item in page["renderedItems"]:
            add_item(slide, item, metrics, colors, chrome)

        notes = slide.notes_slide.notes_text_frame
        notes.text = (
            f"templateId={model['template']['templateId']}; "
            f"fixedImage={fixed_image_path.name}; "
            f"brandLogo={logo_path.name}; "
            f"itemBar={Path(chrome['bar']['resolvedPath']).name}; "
            f"itemNumberBlock={Path(chrome['numberBlock']['resolvedPath']).name}; "
            f"page={page['page']}; href={page['href']}; "
            f"items={len(page['renderedItems'])}; "
            f"compressed={metrics.get('compressed')}"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def main(argv):
    if len(argv) != 3:
        print(
            "用法: python3 scripts/build-agenda-pptx.py <agenda-template.json> <output.pptx>",
            file=sys.stderr,
        )
        return 64
    model_path = Path(argv[1]).resolve()
    output_path = Path(argv[2]).resolve()
    model = json.loads(model_path.read_text(encoding="utf-8"))
    build_pptx(model, model_path, output_path)
    print(f"已生成目录页 PPTX：{output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
