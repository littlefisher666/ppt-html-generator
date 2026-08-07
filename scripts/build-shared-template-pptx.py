#!/usr/bin/env python3
"""由 shared-template.json 生成恒达智控 P3 共享母版 PPTX。"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu

EMU_PER_PX = 7620


def px(value):
    return Emu(round(float(value) * EMU_PER_PX))


def add_rect(slide, x, y, w, h, color=None, transparency=0):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, px(x), px(y), px(w), px(h))
    shape.line.fill.background()
    if color:
        fill = shape.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(color.lstrip("#"))
        fill.transparency = transparency
    else:
        shape.fill.background()
    return shape


def gradient_overlay(shape, start_hex, start_alpha, end_hex, end_alpha, angle=0):
    """写入 a:gradFill；alpha 取 0–100000，angle 为 OOXML 角度（60000=1°）。"""
    sp_pr = shape._element.spPr
    solid = sp_pr.find(qn("a:solidFill"))
    if solid is not None:
        sp_pr.remove(solid)
    no_fill = sp_pr.find(qn("a:noFill"))
    if no_fill is not None:
        sp_pr.remove(no_fill)
    grad = etree.SubElement(sp_pr, qn("a:gradFill"), rotWithShape="1")
    gs_lst = etree.SubElement(grad, qn("a:gsLst"))
    for position, color, alpha in ((0, start_hex, start_alpha), (100000, end_hex, end_alpha)):
        gs = etree.SubElement(gs_lst, qn("a:gs"), pos=str(position))
        srgb = etree.SubElement(gs, qn("a:srgbClr"), val=color.lstrip("#"))
        etree.SubElement(srgb, qn("a:alpha"), val=str(alpha))
    etree.SubElement(grad, qn("a:lin"), ang=str(angle), scaled="1")


def add_picture(slide, image, x, y, w, h):
    slide.shapes.add_picture(str(image), px(x), px(y), width=px(w), height=px(h))


def add_anchor(slide, name, x, y, w, h):
    """无填充无边框锚点；供 P3/P4 读取位置，渲染时不可见。"""
    shape = add_rect(slide, x, y, w, h)
    shape.name = name
    return shape


def build(model, output_path: Path):
    canvas = model["canvas"]
    brand = {key: Path(value) for key, value in model["brand"].items()}
    prs = Presentation()
    prs.slide_width = px(canvas["width"])
    prs.slide_height = px(canvas["height"])
    blank = prs.slide_layouts[6]

    # 封面：品牌原图 + 与 CSS 一致的左→右蓝色渐变遮罩。
    slide = prs.slides.add_slide(blank)
    slide.name = "P3-封面共享模板"
    add_picture(slide, brand["coverBackground"], 0, 0, 1600, 900)
    overlay = add_rect(slide, 0, 0, 1600, 900)
    gradient_overlay(overlay, "#00468F", 84000, "#00468F", 24000, 0)
    add_picture(slide, brand["logoWhite"], 1226.88, 57.57, 331.60, 48.47)
    add_rect(slide, 97.45, 461.40, 519.94, 5, "#44B2FA")
    fade = add_rect(slide, 617.39, 461.40, 620.01, 2.5)
    gradient_overlay(fade, "#BFBFBF", 100000, "#BFBFBF", 0, 0)
    add_rect(slide, 97.45, 640.56, 3.33, 70.87, "#FFFFFF")
    add_anchor(slide, "cover-title-box", 86.67, 233.37, 1395.07, 218.62)
    add_anchor(slide, "cover-subtitle-box", 86.67, 469.25, 1036.76, 70.32)
    add_anchor(slide, "cover-meta-box", 111.71, 635.34, 1125.69, 79.17)

    # 章节：品牌原图 + 左→右蓝色渐变遮罩与可填充内容区锚点。
    slide = prs.slides.add_slide(blank)
    slide.name = "P3-章节共享模板"
    add_picture(slide, brand["sectionBackground"], 0, 0, 1600, 900)
    overlay = add_rect(slide, 0, 0, 1600, 900)
    gradient_overlay(overlay, "#00468F", 86000, "#00468F", 22000, 0)
    add_picture(slide, brand["logoWhite"], 1226.88, 57.57, 331.60, 48.47)
    add_anchor(slide, "section-number-box", 143.35, 206.31, 1003.27, 216.67)
    divider = add_rect(slide, 143.35, 427.0, 622, 4)
    gradient_overlay(divider, "#44B2FA", 100000, "#FFFFFF", 78000, 0)
    add_anchor(slide, "section-title-box", 143.35, 455, 1003.27, 65)
    add_anchor(slide, "section-subtitle-box", 143.35, 532, 1003.27, 45)

    # 内容页：白底、Logo、母版三段线和安全区锚点。
    slide = prs.slides.add_slide(blank)
    slide.name = "P3-内容页共享模板"
    add_rect(slide, 0, 0, 1600, 900, "#FFFFFF")
    # 留出 0.03px 精度余量，避免 OOXML 四舍五入将最右端判为越界。
    fade = add_rect(slide, 277.17, 142.35, 1322.80, 2.5)
    gradient_overlay(fade, "#BFBFBF", 80000, "#BFBFBF", 0, 0)
    add_rect(slide, 157.53, 142.35, 425.20, 5, "#006CB2")
    add_rect(slide, 86.67, 142.35, 70.87, 5, "#44B2FA")
    add_picture(slide, brand["logoColor"], 1240.00, 40.67, 348.17, 53.66)
    add_anchor(slide, "content-title-box", 86.67, 47, 1117.33, 70.87)
    add_anchor(slide, "content-subtitle-box", 86.67, 154, 1425, 35)
    add_anchor(slide, "content-safe-area", 86.67, 207, 1425, 603.67)

    # 封底：品牌原图 + 上→下蓝色渐变遮罩与固定分隔线。
    slide = prs.slides.add_slide(blank)
    slide.name = "P3-封底共享模板"
    add_picture(slide, brand["closingBackground"], 0, 0, 1600, 900)
    overlay = add_rect(slide, 0, 0, 1600, 900)
    gradient_overlay(overlay, "#00468F", 36000, "#002050", 42000, 5400000)
    add_picture(slide, brand["logoWhite"], 1226.88, 57.57, 331.60, 48.47)
    divider = add_rect(slide, 405.45, 365, 789.09, 2)
    gradient_overlay(divider, "#FFFFFF", 0, "#FFFFFF", 90000, 0)
    add_anchor(slide, "closing-title-box", 405.45, 185.39, 789.09, 150)
    add_anchor(slide, "closing-subtitle-box", 405.45, 413, 789.09, 45)
    add_anchor(slide, "closing-org-box", 405.45, 480, 789.09, 45)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def main(argv):
    if len(argv) != 3:
        print("用法: build-shared-template-pptx.py <shared-template.json> <output.pptx>", file=sys.stderr)
        return 64
    model_path = Path(argv[1]).resolve()
    model = json.loads(model_path.read_text(encoding="utf-8"))
    build(model, Path(argv[2]).resolve())
    print(f"已生成 P3 共享模板 PPTX：{Path(argv[2]).resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
