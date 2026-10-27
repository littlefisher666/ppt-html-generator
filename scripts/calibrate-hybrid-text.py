#!/usr/bin/env python3
"""Calibrate LibreOffice text-baseline placement for the P3 hybrid deck.

For every (fontSizePt, lineHeightRatio) combo used by the deck, render an isolated
two-line probe with exact line spacing, then measure where the baseline actually
lands relative to the text-box top. The P3 builder uses the result to place each
text box so that its baseline matches the HTML baseline.

Output: P3/calibration.json

Why exact spacing: LibreOffice multiplies a:spcPct by the font's natural line
height (≈1.205em for Microsoft YaHei), so a CSS line-height of 1.2em becomes
≈1.45em. a:spcPts (exact) reproduces the CSS line box.
"""

from __future__ import annotations

import importlib.util
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

EMU_PER_PX = 7620
CANVAS_WIDTH_PX = 1600
CANVAS_HEIGHT_PX = 900
PROBE_TEXT = "国国国国"
# Chrome ink extents of 国国国国 in Microsoft YaHei, per em (stable across sizes).
INK_ASCENT_RATIO = 0.7806
INK_DESCENT_RATIO = 0.1313
BACKGROUND_GRAY = 0x80


def load_text_writer():
    source = Path(__file__).with_name("pptx-text-writer.py")
    spec = importlib.util.spec_from_file_location("pptx_text_writer", source)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def collect_combos(text_nodes: dict) -> list[tuple[int, float]]:
    combos = set()
    for slide in text_nodes.get("slides") or []:
        for text in slide.get("texts") or []:
            style = text.get("style") or {}
            pt = style.get("fontSizePt")
            ratio = style.get("lineHeightRatio")
            if not pt or not ratio:
                continue
            combos.add((int(pt), round(float(ratio), 4)))
    return sorted(combos)


def build_probe(combos, writer, work_dir: Path) -> tuple[Path, list[dict]]:
    prs = Presentation()
    prs.slide_width = Emu(CANVAS_WIDTH_PX * EMU_PER_PX)
    prs.slide_height = Emu(CANVAS_HEIGHT_PX * EMU_PER_PX)
    blank = prs.slide_layouts[6]
    bands: list[dict] = []

    state = {"slide": None, "y": 0, "slideIndex": 0}

    def new_slide():
        slide = prs.slides.add_slide(blank)
        state["slideIndex"] += 1
        state["slide"] = slide
        state["y"] = 24
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        style = background._element.find(qn("p:style"))
        if style is not None:
            background._element.remove(style)
        background.fill.solid()
        background.fill.fore_color.rgb = RGBColor(BACKGROUND_GRAY, BACKGROUND_GRAY, BACKGROUND_GRAY)
        background.line.fill.background()
        background.name = "probe-background"
        return slide

    for pt, ratio in combos:
        font_px = pt / 0.6
        line_px = font_px * ratio
        pad_top = max(10.0, font_px * 0.08)
        band_height = pad_top + line_px * 2 + max(30.0, font_px * 0.2)
        if state["slide"] is None or state["y"] + band_height > CANVAS_HEIGHT_PX - 10:
            new_slide()
        y = state["y"]
        band = {
            "pt": pt,
            "ratio": ratio,
            "key": f"{pt}|{ratio:g}",
            "lineHeightPx": line_px,
            "slide": state["slideIndex"],
            "x": 40,
            "y": y,
            "boxY": y + pad_top,
            "width": CANVAS_WIDTH_PX - 80,
            "height": band_height,
            "fontPx": font_px,
        }
        spec = {
            "id": f"cal-{pt}-{ratio}",
            "shapeName": f"p3-text:cal-{pt}-{ratio}",
            "box": {"x": band["x"], "y": band["boxY"], "width": band["width"], "height": band_height - pad_top},
            "style": {
                "fontName": "Microsoft YaHei",
                "fontSizePt": pt,
                "bold": False,
                "color": "000000",
                "textAlign": "left",
                "lineHeightRatio": ratio,
                "lineSpacingPx": line_px,
                "letterSpacingPx": 0,
            },
            "text": f"{PROBE_TEXT}\n{PROBE_TEXT}",
            "write": {"wordWrap": False, "useHardLineBreaks": True},
        }
        writer.add_textbox(state["slide"], spec)
        bands.append(band)
        state["y"] = y + band_height + 12

    probe_path = work_dir / "probe.pptx"
    prs.save(str(probe_path))
    return probe_path, bands


def render_probe(pptx: Path, work_dir: Path) -> list[Path]:
    soffice = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if not soffice.is_file():
        binary = shutil.which("soffice") or shutil.which("libreoffice")
        if not binary:
            raise FileNotFoundError("未找到 LibreOffice soffice")
        soffice = Path(binary)
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise FileNotFoundError("未找到 pdftoppm")

    subprocess.run(
        [
            str(soffice), "--headless", "--nologo", "--nolockcheck", "--nodefault",
            "--nofirststartwizard", "--convert-to", "pdf", "--outdir", str(work_dir), str(pptx),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    pdfs = sorted(work_dir.glob("*.pdf"))
    if not pdfs:
        raise RuntimeError("LibreOffice 未生成 PDF")
    prefix = work_dir / "page"
    subprocess.run([pdftoppm, "-png", "-r", "120", str(pdfs[0]), str(prefix)], check=True, capture_output=True)
    return sorted(work_dir.glob("page-*.png"), key=lambda p: int(re.search(r"(\d+)", p.stem).group(1)))


def measure_lines(image, band: dict) -> tuple[tuple[int, int], tuple[int, int]] | None:
    """Ink top/bottom of both probe lines, split at the lightest row between them."""
    line_px = band["lineHeightPx"]
    x0 = int(round(band["x"]))
    x1 = int(round(band["x"] + band["width"]))
    y0 = int(round(band["y"]))
    y1 = min(image.height, int(round(band["y"] + band["height"])))
    if y1 <= y0:
        return None
    crop = image.crop((x0, y0, x1, y1))
    mask = crop.point(lambda value: 255 if value < 110 else 0)
    profile = []
    for row in range(mask.height):
        box = mask.crop((0, row, mask.width, row + 1)).getbbox()
        profile.append(0 if box is None else box[2] - box[0])
    dark_rows = [index for index, value in enumerate(profile) if value > 0]
    if not dark_rows:
        return None
    first, last = dark_rows[0], dark_rows[-1]
    expected = band["boxY"] + line_px - y0
    lo = max(first + 1, int(round(expected - line_px * 0.2)))
    hi = min(last, int(round(expected + line_px * 0.2)))
    if hi <= lo:
        return None
    split = min(range(lo, hi), key=lambda index: (profile[index], index))
    line1_rows = [index for index in dark_rows if index < split]
    line2_rows = [index for index in dark_rows if index >= split]
    if not line1_rows or not line2_rows:
        return None
    return ((line1_rows[0] + y0, line1_rows[-1] + y0), (line2_rows[0] + y0, line2_rows[-1] + y0))


def main() -> int:
    if len(sys.argv) not in (3, 4):
        sys.stderr.write(
            "用法: python3 calibrate-hybrid-text.py <text-nodes.json> <calibration.json> [work-dir]\n"
        )
        return 64

    text_nodes_path = Path(sys.argv[1]).resolve()
    output_path = Path(sys.argv[2]).resolve()
    work_dir = Path(sys.argv[3]).resolve() if len(sys.argv) == 4 else output_path.parent / "calibration"

    text_nodes = json.loads(text_nodes_path.read_text(encoding="utf-8"))
    combos = collect_combos(text_nodes)
    if not combos:
        raise RuntimeError("text-nodes.json 中没有任何字号/行高组合")

    writer = load_text_writer()
    work_dir.mkdir(parents=True, exist_ok=True)
    for old in work_dir.glob("*.pdf"):
        old.unlink()
    for old in work_dir.glob("page-*.png"):
        old.unlink()

    probe_path, bands = build_probe(combos, writer, work_dir)
    pages = render_probe(probe_path, work_dir)

    from PIL import Image

    images = {index + 1: Image.open(path).convert("L") for index, path in enumerate(pages)}
    by_combo: dict[str, dict] = {}
    failures = []
    for band in bands:
        measured = measure_lines(images[band["slide"]], band)
        if not measured:
            failures.append({"combo": band["key"]})
            continue
        line1, line2 = measured
        ink_ascent = INK_ASCENT_RATIO * band["fontPx"]
        ink_descent = INK_DESCENT_RATIO * band["fontPx"]
        baseline1 = ((line1[0] + ink_ascent) + (line1[1] - ink_descent)) / 2
        baseline2 = ((line2[0] + ink_ascent) + (line2[1] - ink_descent)) / 2
        by_combo[band["key"]] = {
            "fontSizePt": band["pt"],
            "lineHeightRatio": band["ratio"],
            "lineHeightPx": round(band["lineHeightPx"], 3),
            "baselineFromBox": round(baseline1 - band["boxY"], 3),
            "baselineGapObserved": round(baseline2 - baseline1, 3),
            "inkTopFromBox": round(line1[0] - band["boxY"], 3),
        }

    if failures:
        raise RuntimeError(f"探针测量失败（未找到两行墨迹）：{json.dumps(failures, ensure_ascii=False)}")

    payload = {
        "stage": "P3",
        "purpose": "text-baseline-calibration",
        "generatedAt": __import__("datetime").datetime.now().isoformat(timespec="seconds"),
        "renderer": "LibreOffice + pdftoppm 120dpi",
        "spacingMode": "exact-spcPts",
        "probe": {"text": PROBE_TEXT, "inkAscentRatio": INK_ASCENT_RATIO, "inkDescentRatio": INK_DESCENT_RATIO},
        "byCombo": by_combo,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "output": str(output_path),
                "combos": len(by_combo),
                "probeSlides": len(pages),
                "maxGapDrift": round(
                    max(
                        abs(item["baselineGapObserved"] - item["lineHeightPx"])
                        for item in by_combo.values()
                    ),
                    3,
                ),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:  # noqa: BLE001
        sys.stderr.write(f"{error}\n")
        raise SystemExit(1)
