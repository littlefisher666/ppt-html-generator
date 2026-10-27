#!/usr/bin/env python3
"""Assemble the P3 hybrid deck: bitmap backgrounds + native editable text boxes.

Backgrounds come from capture-hybrid-backgrounds.mjs (editable glyphs hidden;
vertical-writing / translucent decoration glyphs intentionally kept in the bitmap).
Text specs come from P3/text-nodes.json via build-hybrid-deck.mjs and must not be
hand-authored.
"""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

EMU_PER_PX = 7620
CANVAS_WIDTH_PX = 1600
CANVAS_HEIGHT_PX = 900
PICTURE_NAME = "p3-background"


def load_text_writer():
    """Load the shared writer so assembly and calibration cannot drift."""
    source = Path(__file__).with_name("pptx-text-writer.py")
    spec = importlib.util.spec_from_file_location("pptx_text_writer", source)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main() -> int:
    if len(sys.argv) != 4:
        sys.stderr.write(
            "用法: python3 apply-hybrid-text-plan.py <backgrounds-manifest.json> <text-plan.json> <output.pptx>\n"
        )
        return 64

    manifest_path = Path(sys.argv[1]).resolve()
    plan_path = Path(sys.argv[2]).resolve()
    output_path = Path(sys.argv[3]).resolve()

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    if not plan.get("policy", {}).get("allTextEditable", False):
        raise RuntimeError("text-plan.json 未声明 allTextEditable=true，拒绝组装")

    slides_plan = plan.get("slides") or []
    if len(slides_plan) != len(manifest.get("pages") or []):
        raise RuntimeError("text-plan 页数与 backgrounds manifest 页数不一致")

    writer = load_text_writer()

    prs = Presentation()
    prs.slide_width = Emu(CANVAS_WIDTH_PX * EMU_PER_PX)
    prs.slide_height = Emu(CANVAS_HEIGHT_PX * EMU_PER_PX)
    blank_layout = prs.slide_layouts[6]
    pptx_dir = output_path.parent.parent

    written = 0
    for index, slide_plan in enumerate(slides_plan):
        if slide_plan.get("page") != index + 1:
            raise RuntimeError(
                f"第 {index + 1} 页 page={slide_plan.get('page')} 与页序不一致"
            )
        slide = prs.slides.add_slide(blank_layout)

        background = pptx_dir / str(slide_plan["background"])
        if not background.is_file():
            raise FileNotFoundError(f"缺少背景图: {background}")
        picture = slide.shapes.add_picture(
            str(background), 0, 0, width=prs.slide_width, height=prs.slide_height
        )
        picture.name = PICTURE_NAME
        picture._element.nvPicPr.cNvPr.set("descr", f"P3 无文字背景 · {slide_plan.get('href')}")

        for spec in slide_plan.get("texts") or []:
            if not spec.get("style", {}).get("fontName"):
                raise RuntimeError(f"文本 {spec.get('id')} 缺少 fontName；计划必须来自 text-metrics")
            writer.add_textbox(slide, spec)
            written += 1

        slide.notes_slide.notes_text_frame.text = f"源页面: {slide_plan.get('href')}"

    prs.core_properties.title = str(plan.get("title") or "混合版演示文稿")
    prs.core_properties.subject = "位图背景 + 原生可编辑文字"
    prs.core_properties.comments = "背景来自 P3/capture-hybrid-backgrounds；文字来自 P3/text-nodes.json"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(
        json.dumps(
            {
                "output": str(output_path),
                "pages": len(slides_plan),
                "writtenTexts": written,
                "sizeBytes": output_path.stat().st_size,
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:  # noqa: BLE001
        sys.stderr.write(f"{error}\n")
        raise SystemExit(1)
