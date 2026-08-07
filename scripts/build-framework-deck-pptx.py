#!/usr/bin/env python3
"""Build brand P3 framework-deck.pptx from framework-model.json, with temporary guide-frame overlays."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU_PER_PX = 7620
FONT_UI = "Microsoft YaHei"


def px(value):
    return int(round(float(value) * EMU_PER_PX))


def emu(value):
    return Emu(px(value))


def hex_rgb(value):
    text = value.strip().lstrip("#")
    return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


GUIDE_COLOR = "#FF00A8"
GUIDE_NAME_PREFIX = "guide-frame:"


def add_guide_rect(slide, x, y, w, h, *, line_pt=0.75, name=None):
    """Stroke-only magenta rectangle used as temporary layout guide."""
    if float(w) <= 0 or float(h) <= 0:
        return None
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h)
    )
    shape.fill.background()
    shape.line.color.rgb = hex_rgb(GUIDE_COLOR)
    shape.line.width = Pt(line_pt)
    if name:
        shape.name = name
    return shape


def paint_guide_frames(slide, elements):
    """Overlay P2 guide frames on a framework slide.

    Shape names are prefixed with guide-frame: so finalization can remove them.
    """
    add_guide_rect(slide, 0, 0, 1600, 900, line_pt=1.5, name=f"{GUIDE_NAME_PREFIX}canvas")
    for element in elements or []:
        eid = str(element.get("id") or "el")
        x = float(element.get("x") or 0)
        y = float(element.get("y") or 0)
        width = float(element.get("width") or 0)
        height = float(element.get("height") or 0)
        add_guide_rect(
            slide,
            x,
            y,
            width,
            height,
            line_pt=0.75,
            name=f"{GUIDE_NAME_PREFIX}{eid}",
        )
        if width >= 24 and height >= 14:
            label = add_textbox(
                slide,
                x + 2,
                y + 1,
                min(width - 2, 260),
                14,
                eid,
                font_pt=10,
                color=GUIDE_COLOR,
                bold=False,
                word_wrap=False,
            )
            try:
                label.name = f"{GUIDE_NAME_PREFIX}label:{eid}"
            except Exception:
                pass


def add_rect(slide, x, y, w, h, color=None, transparency=0):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    shape.line.fill.background()
    if color:
        fill = shape.fill
        fill.solid()
        fill.fore_color.rgb = hex_rgb(color)
        fill.transparency = transparency
    else:
        shape.fill.background()
    return shape


def gradient_overlay(shape, start_hex, start_alpha, end_hex, end_alpha, angle=0):
    sp_pr = shape._element.spPr
    solid = sp_pr.find(qn("a:solidFill"))
    if solid is not None:
        sp_pr.remove(solid)
    no_fill = sp_pr.find(qn("a:noFill"))
    if no_fill is not None:
        sp_pr.remove(no_fill)
    grad = etree.SubElement(sp_pr, qn("a:gradFill"), rotWithShape="1")
    gs_lst = etree.SubElement(grad, qn("a:gsLst"))
    for position, color, alpha in ((0, start_hex, start_alpha), (100000, end_hex, end_alpha)):
        gs = etree.SubElement(gs_lst, qn("a:gs"), pos=str(position))
        srgb = etree.SubElement(gs, qn("a:srgbClr"), val=color.lstrip("#"))
        etree.SubElement(srgb, qn("a:alpha"), val=str(alpha))
    etree.SubElement(grad, qn("a:lin"), ang=str(angle), scaled="1")


def add_picture(slide, image, x, y, w, h):
    slide.shapes.add_picture(str(image), emu(x), emu(y), width=emu(w), height=emu(h))


def add_anchor(slide, name, x, y, w, h):
    shape = add_rect(slide, x, y, w, h)
    shape.name = name
    return shape


def zero_text_frame_margins(tf):
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)


def disable_autofit(tf):
    try:
        body_pr = tf._txBody.bodyPr
        for child in list(body_pr):
            if child.tag.endswith("spAutoFit") or child.tag.endswith("normAutofit"):
                body_pr.remove(child)
    except Exception:
        pass


def set_run_font(run, name, size_pt, bold, color):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = hex_rgb(color)
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", name)
    latin = r_pr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(r_pr, qn("a:latin"))
    latin.set("typeface", name)
    r_pr.set("b", "1" if bold else "0")


def add_textbox(slide, x, y, w, h, text, *, font_pt, color="#FFFFFF", bold=False, align="left", anchor="top", font_name=FONT_UI, line_spacing=None, word_wrap=True):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = word_wrap
    zero_text_frame_margins(tf)
    disable_autofit(tf)
    anchor_map = {"bottom": "b", "middle": "ctr", "top": "t"}
    try:
        tf._txBody.bodyPr.set("anchor", anchor_map.get(anchor, "t"))
    except Exception:
        pass
    lines = str(text or "").split("\n") or [""]
    for index, line in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        if line_spacing:
            try:
                pPr = p._p.get_or_add_pPr()
                for child in list(pPr):
                    if child.tag.endswith("lnSpc"):
                        pPr.remove(child)
                lnSpc = etree.SubElement(pPr, qn("a:lnSpc"))
                spcPct = etree.SubElement(lnSpc, qn("a:spcPct"))
                spcPct.set("val", str(int(round(float(line_spacing) * 100000))))
            except Exception:
                pass
        if p.runs:
            p.runs[0].text = ""
        run = p.add_run()
        run.text = line
        set_run_font(run, font_name, float(font_pt), bold, color)
    return box


def resolve_asset_path(model_path: Path, asset: dict, label: str, deck_root: Path) -> Path:
    candidates = []
    if asset.get("path"):
        candidates.append(Path(asset["path"]))
    if asset.get("relativePath"):
        rel = asset["relativePath"]
        candidates.extend([
            model_path.parent / rel,
            deck_root / rel,
            deck_root / "pptx" / rel,
            deck_root / "pptx" / "P3" / rel,
        ])
    for path in candidates:
        if path and path.is_file():
            return path.resolve()
    raise FileNotFoundError(f"missing {label}: {asset}")


def resolve_item_chrome(agenda_model, model_path: Path, item_height, deck_root: Path) -> dict:
    chromes = agenda_model.get("itemChromes") or {}
    chrome = chromes.get(str(item_height))
    if chrome is None:
        for cand_key, cand in chromes.items():
            try:
                if abs(float(cand_key) - float(item_height)) < 0.01:
                    chrome = cand
                    break
            except Exception:
                continue
    if chrome is None:
        raise FileNotFoundError(f"missing agenda item chrome for height={item_height}")
    bar = chrome.get("bar")
    number_block = chrome.get("numberBlock")
    if not bar or not number_block:
        raise ValueError(f"incomplete agenda chrome for height={item_height}")
    return {
        **chrome,
        "bar": {**bar, "resolvedPath": resolve_asset_path(model_path, bar, f"item-bar({item_height})", deck_root)},
        "numberBlock": {**number_block, "resolvedPath": resolve_asset_path(model_path, number_block, f"item-number({item_height})", deck_root)},
    }


def paint_cover_chrome(slide, brand):
    add_picture(slide, brand["coverBackground"], 0, 0, 1600, 900)
    overlay = add_rect(slide, 0, 0, 1600, 900)
    gradient_overlay(overlay, "#00468F", 84000, "#00468F", 24000, 0)
    add_picture(slide, brand["logoWhite"], 1226.88, 57.57, 331.60, 48.47)
    add_rect(slide, 97.45, 461.40, 519.94, 5, "#44B2FA")
    fade = add_rect(slide, 617.39, 461.40, 620.01, 2.5)
    gradient_overlay(fade, "#BFBFBF", 100000, "#BFBFBF", 0, 0)
    add_rect(slide, 97.45, 640.56, 3.33, 70.87, "#FFFFFF")


def paint_section_chrome(slide, brand):
    add_picture(slide, brand["sectionBackground"], 0, 0, 1600, 900)
    overlay = add_rect(slide, 0, 0, 1600, 900)
    gradient_overlay(overlay, "#00468F", 86000, "#00468F", 22000, 0)
    add_picture(slide, brand["logoWhite"], 1226.88, 57.57, 331.60, 48.47)
    divider = add_rect(slide, 143.35, 427.0, 622, 4)
    gradient_overlay(divider, "#44B2FA", 100000, "#FFFFFF", 78000, 0)


def paint_content_chrome(slide, brand):
    add_rect(slide, 0, 0, 1600, 900, "#FFFFFF")
    fade = add_rect(slide, 277.17, 142.35, 1322.80, 2.5)
    gradient_overlay(fade, "#BFBFBF", 80000, "#BFBFBF", 0, 0)
    add_rect(slide, 157.53, 142.35, 425.20, 5, "#006CB2")
    add_rect(slide, 86.67, 142.35, 70.87, 5, "#44B2FA")
    add_picture(slide, brand["logoColor"], 1240.00, 40.67, 348.17, 53.66)
    add_anchor(slide, "content-safe-area", 86.67, 207, 1425, 603.67)


def paint_closing_chrome(slide, brand):
    add_picture(slide, brand["closingBackground"], 0, 0, 1600, 900)
    overlay = add_rect(slide, 0, 0, 1600, 900)
    gradient_overlay(overlay, "#00468F", 36000, "#002050", 42000, 5400000)
    add_picture(slide, brand["logoWhite"], 1226.88, 57.57, 331.60, 48.47)
    divider = add_rect(slide, 405.45, 365, 789.09, 2)
    gradient_overlay(divider, "#FFFFFF", 0, "#FFFFFF", 90000, 0)


def add_agenda_item(slide, item, metrics, colors, chrome):
    bar = chrome["bar"]
    bar_offset = bar.get("placeOffset") or {"x": 0, "y": 0}
    bar_logical = bar.get("logical") or {"width": item["width"], "height": item["height"]}
    add_picture(
        slide,
        bar["resolvedPath"],
        float(item["x"]) + float(bar_offset.get("x", 0)),
        float(item["y"]) + float(bar_offset.get("y", 0)),
        float(bar_logical.get("width", item["width"])),
        float(bar_logical.get("height", item["height"])),
    )
    number_block = chrome["numberBlock"]
    number_offset = number_block.get("placeOffset") or {"x": 0, "y": 0}
    number_logical = number_block.get("logical") or {
        "width": item["numberBox"]["width"],
        "height": item["numberBox"]["height"],
    }
    add_picture(
        slide,
        number_block["resolvedPath"],
        float(item["x"]) + float(number_offset.get("x", 0)),
        float(item["y"]) + float(number_offset.get("y", 0)),
        float(number_logical.get("width", item["numberBox"]["width"])),
        float(number_logical.get("height", item["numberBox"]["height"])),
    )
    nb = item["numberBox"]
    add_textbox(slide, nb["x"], nb["y"], nb["width"], nb["height"], item["number"], font_pt=metrics["numberFontPt"], color=colors["numberColor"], align="center", anchor="middle", font_name=colors["fontFamily"], word_wrap=False)
    tb = item["titleBox"]
    add_textbox(slide, tb["x"], tb["y"], tb["width"], tb["height"], item["title"], font_pt=metrics["titleFontPt"], color=colors["titleColor"], align="left", anchor="middle", font_name=colors["fontFamily"], line_spacing=1.3)


def paint_agenda_page(slide, agenda_model, agenda_page, agenda_model_path: Path, deck_root: Path):
    fixed = agenda_model.get("fixedImage") or {}
    fixed_path = resolve_asset_path(agenda_model_path, fixed, "agenda fixed image", deck_root)
    logical = fixed.get("logical") or {"x": 0, "y": 0, "width": 1600, "height": 900}
    add_picture(slide, fixed_path, logical.get("x", 0), logical.get("y", 0), logical.get("width", 1600), logical.get("height", 900))
    logo = agenda_model.get("brandLogo") or {}
    logo_path = resolve_asset_path(agenda_model_path, logo, "agenda logo", deck_root)
    add_picture(slide, logo_path, logo["x"], logo["y"], logo["width"], logo["height"])
    layout = agenda_page["layout"]
    metrics = layout["metrics"]
    chrome = resolve_item_chrome(agenda_model, agenda_model_path, metrics["itemHeight"], deck_root)
    colors = {
        "titleColor": agenda_model["template"]["itemBase"]["titleColor"],
        "numberColor": agenda_model["template"]["itemBase"]["numberColor"],
        "fontFamily": agenda_model["template"]["itemBase"]["fontFamily"],
    }
    for item in agenda_page["renderedItems"]:
        add_agenda_item(slide, item, metrics, colors, chrome)


def build(model, output_path: Path):
    canvas = model["canvas"]
    brand = {key: Path(value) for key, value in model["brand"].items()}
    deck_root = Path(model["deckRoot"])
    agenda_model = None
    agenda_model_path = None
    agenda_by_page = {}
    if model.get("agendaTemplatePath"):
        agenda_model_path = Path(model["agendaTemplatePath"]).resolve()
        agenda_model = json.loads(agenda_model_path.read_text(encoding="utf-8"))
        for page in agenda_model.get("pages") or []:
            agenda_by_page[int(page["page"])] = page

    include_guide_frames = bool(model.get("includeGuideFrames", True))
    elements_by_page = {}
    guide_frame_by_page = {}
    if include_guide_frames:
        coord_path = model.get("coordinateMapPath")
        if not coord_path:
            raise RuntimeError("includeGuideFrames requires coordinateMapPath in framework-model.json")
        coord_path = Path(coord_path).resolve()
        if not coord_path.is_file():
            raise FileNotFoundError(f"missing coordinate-map: {coord_path}")
        coordinate_map = json.loads(coord_path.read_text(encoding="utf-8"))
        for item in coordinate_map.get("slides") or []:
            page_no = int(item["page"])
            elements_by_page[page_no] = item.get("elements") or []
            if item.get("guideFrame"):
                guide_frame_by_page[page_no] = item["guideFrame"]

    prs = Presentation()
    prs.slide_width = emu(canvas["width"])
    prs.slide_height = emu(canvas["height"])
    blank = prs.slide_layouts[6]

    for page in model["pages"]:
        slide = prs.slides.add_slide(blank)
        layout = page["layout"]
        slide.name = f"P{page['page']:02d}-{layout}"

        if layout == "cover":
            paint_cover_chrome(slide, brand)
            add_textbox(slide, 98.67, 239.37, 1371.07, 212.62, page.get("title") or "", font_pt=40, color="#FFFFFF", bold=True, anchor="bottom", line_spacing=1.2)
            add_textbox(slide, 98.67, 475.25, 1012.76, 64.32, page.get("subtitle") or "", font_pt=24, color="#FFFFFF", anchor="top")
            meta_lines = page.get("metaLines") or []
            if meta_lines:
                add_textbox(slide, 111.71, 635.34, 1125.69, 79.17, "\n".join(meta_lines), font_pt=14, color="#FFFFFF", anchor="top", line_spacing=1.2)
        elif layout == "section":
            paint_section_chrome(slide, brand)
            add_textbox(slide, 143.35, 206.31, 1003.27, 216.67, page.get("sectionNumber") or "", font_pt=130, color="#FFFFFF", anchor="top", word_wrap=False)
            add_textbox(slide, 143.35, 455, 1003.27, 65, page.get("title") or "", font_pt=32, color="#FFFFFF", bold=True, anchor="top")
            if page.get("subtitle"):
                add_textbox(slide, 143.35, 532, 1003.27, 45, page.get("subtitle") or "", font_pt=18, color="#FFFFFF", anchor="top")
        elif layout == "content":
            paint_content_chrome(slide, brand)
            add_textbox(slide, 86.67, 47, 1117.33, 70.87, page.get("title") or "", font_pt=32, color="#000000", bold=True, anchor="middle", word_wrap=False)
            if page.get("subtitle"):
                add_textbox(slide, 86.67, 154, 1425, 35, page.get("subtitle") or "", font_pt=14, color="#596B82", anchor="middle", word_wrap=False)
        elif layout == "closing":
            paint_closing_chrome(slide, brand)
            add_textbox(slide, 405.45, 185.39, 789.09, 150, page.get("title") or "", font_pt=80, color="#FFFFFF", bold=True, align="center", anchor="middle")
            if page.get("subtitle"):
                add_textbox(slide, 405.45, 413, 789.09, 45, page.get("subtitle") or "", font_pt=18, color="#FFFFFF", align="center", anchor="middle")
            if page.get("org"):
                add_textbox(slide, 405.45, 480, 789.09, 45, page.get("org") or "", font_pt=18, color="#FFFFFF", align="center", anchor="middle")
        elif layout == "agenda":
            if agenda_model is None:
                raise RuntimeError("agenda pages present but agenda-template.json missing")
            agenda_page = agenda_by_page.get(int(page["page"]))
            if agenda_page is None:
                raise RuntimeError(f"agenda-template.json missing page {page['page']}")
            paint_agenda_page(slide, agenda_model, agenda_page, agenda_model_path, deck_root)
        else:
            raise RuntimeError(f"unsupported layout: {layout}")

        if include_guide_frames:
            paint_guide_frames(slide, elements_by_page.get(int(page["page"]), []))

        notes = slide.notes_slide.notes_text_frame
        guide_ref = guide_frame_by_page.get(int(page["page"])) or page.get("guideFrame") or ""
        notes.text = (
            f"page={page['page']}; layout={layout}; href={page.get('href')}; "
            f"title={page.get('title') or ''}; guideFrame={guide_ref}; "
            f"guideShapes={'yes' if include_guide_frames else 'no'}"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def main(argv):
    if len(argv) != 3:
        print("Usage: build-framework-deck-pptx.py <framework-model.json> <output.pptx>", file=sys.stderr)
        return 64
    model_path = Path(argv[1]).resolve()
    output_path = Path(argv[2]).resolve()
    model = json.loads(model_path.read_text(encoding="utf-8"))
    build(model, output_path)
    print(f"generated framework pptx: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
